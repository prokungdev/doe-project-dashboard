#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
extract_details.py
Extracts detailed activity breakdown (rows 8-194) for all 124 units from:
data/1.จัดสรร  ณ 5 ต.ค. ตัดแผนงานพื้นฐานออก_20261007.xlsx
Outputs: dashboard_detail.json
"""
import openpyxl, json, os, sys

def extract():
    sys.stdout.reconfigure(encoding='utf-8')
    base_dir = os.path.dirname(os.path.abspath(__file__))
    excel_path = os.path.join(base_dir, 'data', '1.จัดสรร  ณ 5 ต.ค. ตัดแผนงานพื้นฐานออก_20261007.xlsx')
    
    print(f"Loading {excel_path}...")
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    sheets = [s for s in wb.sheetnames if s != 'Recheck']
    ws_ref = wb['สระบุรี']
    
    raw_rows = []
    for r in range(8, 195):
        c1 = ws_ref.cell(r, 1).value
        c2 = ws_ref.cell(r, 2).value
        raw_rows.append({
            'r': r,
            'c1': str(c1).strip() if c1 is not None else '',
            'c2': str(c2).strip() if c2 is not None else ''
        })
    
    # Check which rows have data in any sheet
    row_has_target = {r: False for r in range(8, 195)}
    row_units = {r: '' for r in range(8, 195)}
    
    for s in sheets:
        ws = wb[s]
        for r in range(8, 195):
            u = ws.cell(r, 2).value
            if u and str(u).strip() and not row_units[r]:
                row_units[r] = str(u).strip()
            v = ws.cell(r, 3).value
            if v not in (None, '', 0):
                row_has_target[r] = True
    
    # Build clean template
    template = []
    current_plan = ""
    current_proj = ""
    current_proj_num = 0
    
    for idx, item in enumerate(raw_rows):
        r = item['r']
        c1 = item['c1']
        c2 = item['c2'] if item['c2'] else row_units[r]
        
        level = 3
        row_type = 'sub_act'
        code = ''
        
        if 'แผนงานยุทธศาสตร์' in c1:
            level = 0
            row_type = 'plan'
            current_plan = c1
            code = 'แผนงาน'
        elif any(c1.startswith(f'{i}. ') for i in range(1, 12)):
            level = 1
            row_type = 'project'
            current_proj = c1
            try:
                current_proj_num = int(c1.split('.')[0])
                code = f'โครงการ {current_proj_num}'
            except:
                pass
        elif any(c1.startswith(f'{current_proj_num}.{j} ') or f'{current_proj_num}.{j} กิจกรรม' in c1 for j in range(1, 10)):
            level = 2
            row_type = 'main_act'
            code = c1.split(' ')[0].replace(':', '')
        elif any(c1.startswith(f'{current_proj_num}.{j}.{k}') or f'{current_proj_num}.{j}.{k}' in c1 for j in range(1, 10) for k in range(1, 15)):
            level = 3
            row_type = 'sub_act'
            code = c1.split(' ')[0].replace(':', '')
        elif any(c1.startswith(f'{i})') or f' {i})' in c1 for i in range(1, 10)):
            level = 4
            row_type = 'item'
            code = c1.split(')')[0].strip() + ')'
        elif c2 != '':
            level = 3.5
            row_type = 'sub_unit'
            code = '↳'
        else:
            level = 3.5
            row_type = 'note'
            code = '•'
    
        disp_name = c1
        if not disp_name and c2:
            disp_name = f"(หน่วย: {c2})"
        elif c2 and row_type == 'sub_unit' and not any(disp_name.startswith(p) for p in ['1.', '2.', '3.', '4.', '5.', '6.', '7.', '8.', '9.', '10.', '11.']):
            disp_name = f"↳ {disp_name}" if not disp_name.startswith('↳') else disp_name
    
        template.append({
            'r': r,
            'level': level,
            'type': row_type,
            'code': code,
            'name': disp_name,
            'raw_name': c1,
            'unit': c2,
            'plan': current_plan,
            'proj': current_proj_num,
            'has_any_target': row_has_target[r]
        })
    
    # Extract matrix for each sheet
    units_data = {}
    for s in sheets:
        ws = wb[s]
        unit_rows = {}
        for r in range(8, 195):
            vals = [ws.cell(r, c).value for c in range(3, 20)]
            cleaned = []
            has_val = False
            for v in vals:
                if v is None or v == '':
                    cleaned.append(0)
                elif isinstance(v, (int, float)):
                    iv = int(v) if v == int(v) else round(v, 2)
                    cleaned.append(iv)
                    if iv != 0:
                        has_val = True
                else:
                    cleaned.append(0)
            if has_val:
                unit_rows[str(r)] = cleaned
        units_data[s] = unit_rows
    
    out_file = os.path.join(base_dir, 'dashboard_detail.json')
    with open(out_file, 'w', encoding='utf-8') as f:
        json.dump({'template': template, 'units': units_data}, f, ensure_ascii=False)
    
    print(f"Detail data successfully saved to: {out_file}")
    print(f"Size: {os.path.getsize(out_file) / 1024:.1f} KB")

if __name__ == '__main__':
    extract()
