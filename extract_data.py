#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
extract_data.py
Extracts project-level target allocations (11 projects) for all 124 units from:
data/1.จัดสรร  ณ 5 ต.ค. ตัดแผนงานพื้นฐานออก_20261007.xlsx
Outputs: dashboard_data.json
"""
import openpyxl
import json
import os
import sys

def extract():
    sys.stdout.reconfigure(encoding='utf-8')
    base_dir = os.path.dirname(os.path.abspath(__file__))
    excel_path = os.path.join(base_dir, 'data', '1.จัดสรร  ณ 5 ต.ค. ตัดแผนงานพื้นฐานออก_20261007.xlsx')
    
    print(f"Loading {excel_path}...")
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    sheets = [s for s in wb.sheetnames if s != 'Recheck']
    
    # 11 project rows
    proj_rows = [9, 18, 36, 43, 52, 58, 65, 106, 156, 173, 180]
    
    # Read existing dashboard_data.json for metadata (project names, unit categories & clean names)
    existing_p = os.path.join(base_dir, 'dashboard_data.json')
    with open(existing_p, 'r', encoding='utf-8') as f:
        existing = json.load(f)
    
    projects_def = existing['projects']
    unit_meta = {u['sheet']: {'name': u['name'], 'category': u['category']} for u in existing['units']}
    
    units_data = []
    
    for s in sheets:
        ws = wb[s]
        meta = unit_meta.get(s, {
            'name': str(ws.cell(2, 1).value or s).strip(),
            'category': 'จังหวัด' if 'จังหวัด' in s or len(s) <= 10 else 'ส่วนกลาง'
        })
        
        proj_list = []
        for p_idx, r in enumerate(proj_rows, 1):
            p_def = projects_def[p_idx - 1]
            p_name = p_def['name']
            
            unit_str = ws.cell(r, 2).value
            unit_str = str(unit_str).strip() if unit_str else 'คน'
            if p_idx == 11:
                unit_str = 'ระบบ'
            
            def get_val(c):
                v = ws.cell(r, c).value
                if v is None or v == '':
                    return 0
                try:
                    return int(v) if float(v) == int(float(v)) else round(float(v), 2)
                except:
                    return 0

            target = get_val(3)
            q1 = {'oct': get_val(4), 'nov': get_val(5), 'dec': get_val(6), 'total': get_val(7)}
            q2 = {'jan': get_val(8), 'feb': get_val(9), 'mar': get_val(10), 'total': get_val(11)}
            q3 = {'apr': get_val(12), 'may': get_val(13), 'jun': get_val(14), 'total': get_val(15)}
            q4 = {'jul': get_val(16), 'aug': get_val(17), 'sep': get_val(18), 'total': get_val(19)}
            
            proj_list.append({
                'num': p_idx,
                'name': p_name,
                'unit': unit_str,
                'target': target,
                'q1': q1,
                'q2': q2,
                'q3': q3,
                'q4': q4
            })
            
        units_data.append({
            'sheet': s,
            'name': meta['name'],
            'category': meta['category'],
            'projects': proj_list
        })
    
    out_file = os.path.join(base_dir, 'dashboard_data.json')
    with open(out_file, 'w', encoding='utf-8') as f:
        json.dump({'projects': projects_def, 'units': units_data}, f, ensure_ascii=False)
        
    print(f"Project-level data successfully saved to: {out_file}")

if __name__ == '__main__':
    extract()
